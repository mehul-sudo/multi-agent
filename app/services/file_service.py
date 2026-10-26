"""File storage and document content extraction service."""

import csv
import os
import uuid
from typing import Any, Dict, List, Optional, Tuple
from werkzeug.utils import secure_filename

from app.extensions import db
from app.models.uploaded_file import UploadedFile

ALLOWED_EXTENSIONS = {
    "csv", "xlsx", "xls", "pdf", "txt", "md", "json", "doc", "docx", "pptx",
    "py", "js", "ts", "jsx", "tsx", "html", "css", "sql", "sh", "yaml", "yml", "log"
}
MAX_FILE_SIZE = 15 * 1024 * 1024  # 15 MB


class FileService:
    """Enterprise file upload, access control, and document extraction service."""

    @staticmethod
    def is_allowed_extension(filename: str) -> bool:
        if "." not in filename:
            return False
        ext = filename.rsplit(".", 1)[1].lower()
        return ext in ALLOWED_EXTENSIONS

    @classmethod
    def save_uploaded_file(
        cls,
        file_storage,
        upload_dir: str,
        user_id: Optional[int] = None,
    ) -> Tuple[Optional[UploadedFile], Optional[str]]:
        """Persist an uploaded file securely with DB metadata and ownership."""
        if not file_storage or not file_storage.filename:
            return None, "No file provided."

        orig_filename = secure_filename(file_storage.filename)
        if not cls.is_allowed_extension(orig_filename):
            return None, f"File type not allowed. Supported types: {', '.join(sorted(ALLOWED_EXTENSIONS))}"

        os.makedirs(upload_dir, exist_ok=True)

        file_storage.seek(0, os.SEEK_END)
        size_bytes = file_storage.tell()
        if size_bytes > MAX_FILE_SIZE:
            return None, f"File exceeds maximum allowed size ({MAX_FILE_SIZE // (1024 * 1024)}MB)."
        file_storage.seek(0)

        file_id = str(uuid.uuid4())
        ext = orig_filename.rsplit(".", 1)[1].lower()
        saved_filename = f"{file_id}.{ext}"
        storage_path = os.path.abspath(os.path.join(upload_dir, saved_filename))

        # Path traversal guard
        if not storage_path.startswith(os.path.abspath(upload_dir)):
            return None, "Invalid file destination path."

        file_storage.save(storage_path)

        # Build initial preview
        preview = cls._generate_preview(storage_path, ext)

        uploaded = UploadedFile(
            public_id=file_id,
            filename=orig_filename,
            extension=ext,
            mime_type=getattr(file_storage, "content_type", None) or f"application/{ext}",
            size_bytes=size_bytes,
            storage_path=storage_path,
            owner_id=user_id,
            metadata_json={"preview": preview},
        )
        db.session.add(uploaded)
        db.session.commit()

        return uploaded, None

    @classmethod
    def get_file_by_id(
        cls, file_id: str, user_id: Optional[int] = None
    ) -> Tuple[Optional[UploadedFile], Optional[str]]:
        """Retrieve verified uploaded file with ownership verification."""
        file_record = UploadedFile.query.filter_by(public_id=file_id).first()
        if not file_record:
            return None, f"File with ID '{file_id}' not found."

        if user_id is not None and file_record.owner_id is not None:
            if file_record.owner_id != user_id:
                return None, "Unauthorized: You do not have permission to access this file."

        if not os.path.exists(file_record.storage_path):
            return None, "File data does not exist on storage."

        return file_record, None

    @classmethod
    def extract_document_content(
        cls, file_path: str, filename: Optional[str] = None, user_id: Optional[int] = None
    ) -> Tuple[List[Dict[str, Any]], Optional[str]]:
        """
        Extract text chunks with precise source metadata (page, slide, paragraph)
        for PDF, DOCX, PPTX, TXT, MD, CSV, and XLSX.
        Detects scanned PDFs with zero text.
        """
        if not os.path.exists(file_path):
            return [], f"File '{file_path}' does not exist."

        name = filename or os.path.basename(file_path)
        ext = name.rsplit(".", 1)[1].lower() if "." in name else ""

        documents: List[Dict[str, Any]] = []

        # 1. PDF EXTRACTION
        if ext == "pdf":
            extracted_pages = []
            try:
                import pdfplumber
                with pdfplumber.open(file_path) as pdf:
                    for idx, page in enumerate(pdf.pages):
                        txt = page.extract_text() or ""
                        if txt.strip():
                            extracted_pages.append((idx + 1, txt.strip()))
            except Exception:
                try:
                    import pypdf
                    reader = pypdf.PdfReader(file_path)
                    for idx, page in enumerate(reader.pages):
                        txt = page.extract_text() or ""
                        if txt.strip():
                            extracted_pages.append((idx + 1, txt.strip()))
                except Exception as e:
                    # Fallback: check if the file is actually a plaintext, CSV, or text-encoded file with .pdf extension
                    try:
                        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                            raw_content = f.read()
                        # If it doesn't start with PDF magic bytes and contains text/csv lines
                        if raw_content and not raw_content.startswith("%PDF-") and any(c.isalnum() for c in raw_content[:200]):
                            documents.append({
                                "id": f"{name}_content",
                                "name": f"{name} (Extracted Text Content)",
                                "text": raw_content[:50000],
                                "page": 1,
                                "owner_id": user_id,
                            })
                            return documents, None
                    except Exception:
                        pass
                    return [], f"Failed to parse PDF document: {str(e)}"

            if not extracted_pages:
                return [], (
                    "Scanned image PDF contains no extractable text. "
                    "Optical Character Recognition (OCR) is not currently installed on the server to extract text from image-only pages. "
                    "Please provide a text-based searchable PDF or plain text document."
                )

            for p_num, p_text in extracted_pages:
                documents.append({
                    "id": f"{name}_p{p_num}",
                    "name": f"{name} (Page {p_num})",
                    "text": p_text,
                    "page": p_num,
                    "owner_id": user_id,
                })
            return documents, None

        # 2. DOCX EXTRACTION
        elif ext in ("docx", "doc"):
            try:
                import docx
                doc = docx.Document(file_path)
                p_items = []
                current_section = "General"

                for p_idx, p in enumerate(doc.paragraphs):
                    text = p.text.strip()
                    if not text:
                        continue
                    if p.style.name.startswith("Heading"):
                        current_section = text
                    p_items.append((p_idx + 1, current_section, text))

                for t_idx, table in enumerate(doc.tables):
                    table_rows = []
                    for row in table.rows:
                        row_cells = [c.text.strip() for c in row.cells if c.text.strip()]
                        if row_cells:
                            table_rows.append(" | ".join(row_cells))
                    if table_rows:
                        p_items.append((
                            len(p_items) + 1,
                            f"Table {t_idx+1}",
                            "\n".join(table_rows),
                        ))

                if not p_items:
                    return [], "DOCX document contains no extractable text."

                # Group paragraphs into logical units
                batch = []
                batch_text_len = 0
                sec = current_section
                start_p = 1

                for p_num, p_sec, p_text in p_items:
                    if batch and (batch_text_len + len(p_text) > 1200 or p_sec != sec):
                        combined_text = "\n\n".join(batch)
                        documents.append({
                            "id": f"{name}_para_{start_p}_{p_num-1}",
                            "name": f"{name} ({sec}, Paras {start_p}-{p_num-1})",
                            "text": combined_text,
                            "paragraph": start_p,
                            "section": sec,
                            "owner_id": user_id,
                        })
                        batch = []
                        batch_text_len = 0
                        start_p = p_num

                    sec = p_sec
                    batch.append(p_text)
                    batch_text_len += len(p_text)

                if batch:
                    documents.append({
                        "id": f"{name}_para_{start_p}_{len(p_items)}",
                        "name": f"{name} ({sec}, Paras {start_p}-{len(p_items)})",
                        "text": "\n\n".join(batch),
                        "paragraph": start_p,
                        "section": sec,
                        "owner_id": user_id,
                    })

                return documents, None

            except Exception as e:
                return [], f"Failed to parse DOCX document: {str(e)}"

        # 3. PPTX EXTRACTION
        elif ext == "pptx":
            try:
                import pptx
                prs = pptx.Presentation(file_path)
                slide_items = []

                def _extract_shape_text(shape) -> List[str]:
                    texts = []
                    if shape.has_text_frame:
                        for paragraph in shape.text_frame.paragraphs:
                            t = paragraph.text.strip()
                            if t:
                                texts.append(t)
                    if shape.has_table:
                        for row in shape.table.rows:
                            row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                            if row_cells:
                                texts.append(" | ".join(row_cells))
                    if hasattr(shape, "shapes"):
                        for sub_shape in shape.shapes:
                            texts.extend(_extract_shape_text(sub_shape))
                    return texts

                for s_idx, slide in enumerate(prs.slides):
                    slide_num = s_idx + 1
                    slide_title = ""

                    try:
                        if slide.shapes.title and slide.shapes.title.text.strip():
                            slide_title = slide.shapes.title.text.strip()
                    except Exception:
                        pass

                    slide_texts = []
                    for shape in slide.shapes:
                        is_title_shape = False
                        try:
                            if slide.shapes.title and shape == slide.shapes.title:
                                is_title_shape = True
                        except Exception:
                            pass

                        extracted = _extract_shape_text(shape)
                        if is_title_shape and not slide_title and extracted:
                            slide_title = extracted[0]
                        elif not is_title_shape:
                            slide_texts.extend(extracted)
                        elif is_title_shape and len(extracted) > 1:
                            slide_texts.extend(extracted[1:])

                    if not slide_title and slide_texts:
                        slide_title = slide_texts[0]

                    if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                        notes_txt = slide.notes_slide.notes_text_frame.text.strip()
                        if notes_txt:
                            slide_texts.append(f"Speaker Notes: {notes_txt}")

                    body_text = "\n".join(slide_texts)
                    content_parts = []
                    if slide_title:
                        content_parts.append(f"Title: {slide_title}")
                    if body_text:
                        content_parts.append(body_text)

                    full_slide_text = "\n".join(content_parts).strip()
                    if full_slide_text:
                        slide_items.append((slide_num, slide_title or f"Slide {slide_num}", full_slide_text))

                if not slide_items:
                    return [], f"PPTX presentation '{name}' contains no extractable text across {len(prs.slides)} slides."

                for s_num, s_title, s_content in slide_items:
                    documents.append({
                        "id": f"{name}_slide_{s_num}",
                        "name": f"{name} (Slide {s_num}: {s_title[:40]})",
                        "text": s_content,
                        "slide": s_num,
                        "slide_title": s_title,
                        "owner_id": user_id,
                    })

                return documents, None

            except Exception as e:
                return [], f"Failed to parse PPTX presentation: {str(e)}"

        # 4. PLAIN TEXT, MARKDOWN & SOURCE CODE
        elif ext in ("txt", "md", "json", "log", "py", "js", "ts", "jsx", "tsx", "html", "css", "sql", "sh", "yaml", "yml"):
            try:
                with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                    lines = f.readlines()

                if not lines or not any(l.strip() for l in lines):
                    return [], f"File '{name}' is empty."

                chunk_lines = []
                start_l = 1
                for idx, line in enumerate(lines):
                    chunk_lines.append(line)
                    if len(chunk_lines) >= 40:
                        txt_block = "".join(chunk_lines).strip()
                        if txt_block:
                            documents.append({
                                "id": f"{name}_lines_{start_l}_{idx+1}",
                                "name": f"{name} (Lines {start_l}-{idx+1})",
                                "text": txt_block,
                                "line_start": start_l,
                                "line_end": idx + 1,
                                "owner_id": user_id,
                            })
                        chunk_lines = []
                        start_l = idx + 2

                if chunk_lines:
                    txt_block = "".join(chunk_lines).strip()
                    if txt_block:
                        documents.append({
                            "id": f"{name}_lines_{start_l}_{len(lines)}",
                            "name": f"{name} (Lines {start_l}-{len(lines)})",
                            "text": txt_block,
                            "line_start": start_l,
                            "line_end": len(lines),
                            "owner_id": user_id,
                        })

                return documents, None
            except Exception as e:
                return [], f"Failed to read text file: {str(e)}"

        # 5. CSV & TABULAR DATA
        elif ext in ("csv", "xlsx", "xls"):
            try:
                import pandas as pd
                if ext == "csv":
                    df = pd.read_csv(file_path)
                else:
                    df = pd.read_excel(file_path)

                if df.empty:
                    return [], f"Data file '{name}' contains no rows."

                columns = list(df.columns)
                num_rows = len(df)
                summary_text = (
                    f"Dataset: {name}\nColumns ({len(columns)}): {', '.join(map(str, columns))}\n"
                    f"Total Rows: {num_rows}\n\nSummary Statistics:\n{df.describe(include='all').to_string()}"
                )
                documents.append({
                    "id": f"{name}_summary",
                    "name": f"{name} (Dataset Overview)",
                    "text": summary_text,
                    "owner_id": user_id,
                })

                step = 25
                for i in range(0, min(num_rows, 200), step):
                    subset = df.iloc[i : i + step]
                    documents.append({
                        "id": f"{name}_rows_{i+1}_{i+len(subset)}",
                        "name": f"{name} (Rows {i+1} to {i+len(subset)})",
                        "text": subset.to_string(),
                        "owner_id": user_id,
                    })

                return documents, None
            except Exception as e:
                return [], f"Failed to parse tabular dataset: {str(e)}"

        return [], f"Unsupported document format: .{ext}"

    @classmethod
    def _generate_preview(cls, file_path: str, ext: str) -> Dict[str, Any]:
        """Generate safe, small preview metadata for client response."""
        try:
            if ext == "csv":
                with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                    reader = csv.reader(f)
                    rows = [next(reader, None) for _ in range(5)]
                    rows = [r for r in rows if r is not None]
                    if rows:
                        return {"columns": rows[0], "rows": rows[1:]}
            elif ext == "pdf":
                import pdfplumber
                with pdfplumber.open(file_path) as pdf:
                    first_page = pdf.pages[0].extract_text() if pdf.pages else ""
                    return {"page_count": len(pdf.pages), "sample": (first_page or "")[:300]}
            elif ext == "pptx":
                import pptx
                prs = pptx.Presentation(file_path)
                first_slide_text = ""
                if prs.slides:
                    for s in prs.slides[0].shapes:
                        if s.has_text_frame and s.text_frame.text.strip():
                            first_slide_text += s.text_frame.text.strip() + " "
                return {"slide_count": len(prs.slides), "sample": first_slide_text.strip()[:300]}
            elif ext in ("docx", "doc"):
                import docx
                doc = docx.Document(file_path)
                sample = " ".join([p.text.strip() for p in doc.paragraphs if p.text.strip()][:3])
                return {"paragraph_count": len(doc.paragraphs), "sample": sample[:300]}
            elif ext in ("txt", "md"):
                with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                    snippet = f.read(500)
                    return {"sample": snippet}
        except Exception:
            pass
        return {}
